from pathlib import Path
import json
import re
import sys
import importlib.util
from collections import Counter
from docx.enum.section import WD_SECTION_START

sys.stdout.reconfigure(encoding='utf-8')
工作目录 = Path(__file__).resolve().parent
根目录 = 工作目录.parent
英语目录 = 根目录 / '朵朵/6年级/记录/英语'
原稿目录 = 英语目录 / '原稿'
输出目录 = 英语目录 / '错题'
规格 = importlib.util.spec_from_file_location('旧工具', 工作目录 / '整理三份英语错题集.py')
旧工具 = importlib.util.module_from_spec(规格)
规格.loader.exec_module(旧工具)
排版 = 旧工具.排版


def 读取项目(卷号):
    路径 = 工作目录 / f'{卷号}蓝笔错题核对.json'
    if 路径.exists():
        数据 = json.loads(路径.read_text(encoding='utf-8-sig'))
        项目 = 数据.get('项目', [])
    else:
        数据 = json.loads((工作目录 / f'{卷号}错题核对.json').read_text(encoding='utf-8-sig'))
        if 卷号 == 'L008':
            项目 = 数据.get('基础', [])[:]
            for 阅读 in 数据.get('阅读', []):
                项目 += [dict(x, 大题='阅读' + 阅读['组']) for x in 阅读.get('错题', [])]
        else:
            项目 = 数据
    结果 = {}
    for 题 in 项目:
        状态 = 题.get('状态', '')
        if 状态 in ('做错', '错题', '已订正', '未作答', '未完成', '错题（照片已涂改）'):
            大题 = str(题['大题'])
            m = re.search(r'阅读训练\s*([A-E])|能力提升\s*([A-E])|^([A-E])[\s　]', 大题)
            if m:
                大题 = '阅读' + next(x for x in m.groups() if x)
            elif '、' in 大题:
                大题 = 大题.split('、')[0]
            结果[(大题, int(题['题号']))] = dict(题, 大题=大题)
    return 结果


def 清理前文文本(文本):
    """删除双栏题目表头，保留题目前的说明或词语表。"""
    结果 = []
    for 行 in 文本.splitlines():
        行 = 行.strip()
        if not 行 or 行.strip('|').strip() == '' or ('题目' in 行 and '**' not in 行):
            continue
        if re.fullmatch(r'\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)+\|?', 行):
            continue
        结果.append(行)
    return '\n'.join(结果).strip()


def 分解原稿(卷号):
    文本 = (原稿目录 / f'2026-09-16-{卷号}英语试题.md').read_text(encoding='utf-8-sig')
    if 卷号 == 'L001':
        文本 = 文本.replace('Tommy is good at', 'Tommy is a good')
        文本 = 文本.replace('(visit) the city next week', '(go) the city next week')
        文本 = 文本.replace('the Art Museum', 'the Car Museum')
    大节 = ''
    小节 = None
    各节 = []
    for 行 in 文本.splitlines():
        if 行.startswith('## '):
            大节 = 行[3:]
            小节 = None
        elif 行.startswith('### '):
            标题 = 行[4:]
            if 大节 == '基础练习':
                键 = 标题.split('、')[0]
            else:
                m = re.match(r'([A-E])(?:\s|　|$)', 标题)
                if not m:
                    m = re.search(r'阅读训练\s*([A-E])|能力提升\s*([A-E])', 标题)
                键 = '阅读' + (m.group(1) if m and m.group(1) else (m.group(2) if m else 'A'))
            小节 = {'大节': 大节, '标题': 标题, '键': 键, '正文': []}
            各节.append(小节)
        elif 小节 is not None and 行 != '---':
            小节['正文'].append(行)
    def 清理题目文本(文本):
        """去掉双栏表格切分后残留的边界，只保留题目文字。"""
        结果 = []
        for 行 in 文本.splitlines():
            行 = 行.strip()
            if not 行 or re.fullmatch(r'\|?\s*:?-+:?\s*(\|\s*:?-+:?\s*)+\|?', 行):
                continue
            if 行.startswith('|') or 行.endswith('|'):
                行 = 行.strip('|').strip()
            if 行:
                结果.append(行)
        return '\n'.join(结果).replace('<br>', '\n').strip()

    for 节 in 各节:
        正文 = '\n'.join(节['正文']).strip()
        if 卷号 == 'L009' and 节['键'] == '一':
            题目 = {}
            for 行 in 正文.splitlines()[2:]:
                for 格 in 行.strip('|').split('|'):
                    m = re.search(r'\*\*(\d+)\.', 格)
                    if m:
                        题目[int(m.group(1))] = 格.strip()
            节['前文'] = ''
            节['题目'] = 题目
        else:
            划分 = list(re.finditer(r'\*\*(\d+)\.\*\*', 正文))
            节['前文'] = 正文[:划分[0].start()].strip() if 划分 else 正文
            节['题目'] = {
                int(m.group(1)): 清理题目文本(
                    正文[m.start():划分[i+1].start() if i+1 < len(划分) else len(正文)]
                )
                for i, m in enumerate(划分)
            }
    return 各节


def 反馈(题):
    原 = str(题.get('原答案', '未作答'))
    if 题.get('状态') == '未作答' and not 原.startswith('未作答'):
        原 = '未作答；' + 原
    蓝 = str(题.get('蓝笔订正', '') or '')
    项 = [f'**朵朵原答案：** {原}']
    if 蓝:
        项.append(f'**蓝笔订正：** {蓝}')
    项 += ['**订正：**', '＿＿' * 24, '**心得：**', '＿＿' * 24, '＿＿' * 24]
    return '\n\n'.join(项)


def 生成卷Markdown(卷号):
    项目 = 读取项目(卷号)
    各节 = 分解原稿(卷号)
    内容 = [f'# 六年级第一学期英语练习 {卷号}']
    当前大节 = ''
    for 节 in 各节:
        本节 = {号: 题 for (键, 号), 题 in 项目.items() if 键 == 节['键']}
        if not 本节:
            continue
        if 当前大节 != 节['大节']:
            内容.append('## ' + 节['大节'])
            当前大节 = 节['大节']
        内容.append('### ' + 节['标题'])
        阅读 = 节['键'].startswith('阅读')
        前文 = 清理前文文本(节['前文'])
        if 卷号 == 'L010' and 节['键'] == '三':
            # 选项表与对话必须分成两个 Markdown 块，避免 Word 将对话误识别为表格行。
            前文 = re.sub(r'\n(?=A: Hi,)', '\n\n', 前文)
        if 前文:
            内容.append(前文)
        if 阅读 and 节['题目']:
            号码 = list(节['题目'])
        elif not 节['题目']:
            号码 = sorted(本节)
        else:
            号码 = sorted(本节)
        if not 节['题目'] and 卷号 == 'L010' and 节['键'] == '三':
            内容.append('**5.** ＿＿＿＿＿＿＿＿')
            内容.append(反馈(本节[5]))
            continue
        for 号 in 号码:
            if 号 not in 节['题目']:
                内容.append(f'**{号}.** ＿＿＿＿＿＿＿＿')
            else:
                内容.append(节['题目'][号])
            if 号 in 本节:
                内容.append(反馈(本节[号]))
    输出 = 输出目录 / f'2026-09-17-{卷号}蓝笔错题集.md'
    输出.write_text('\n\n'.join(内容) + '\n', encoding='utf-8')
    return 输出, len(项目), dict(Counter(x.get('状态') for x in 项目.values()))


def 排版正文(文档, 文本, 卷号):
    # 原稿中部分表格与说明文字之间没有空行，先按连续竖线行切开，避免整段被当作普通文字。
    表格分隔 = []
    表格中 = False
    for 行 in 文本.replace('\r\n', '\n').split('\n'):
        是表格行 = 行.strip().startswith('|')
        if 是表格行 and not 表格中 and 表格分隔 and 表格分隔[-1] != '':
            表格分隔.append('')
        if not 是表格行 and 表格中 and 行.strip():
            表格分隔.append('')
        表格分隔.append(行)
        表格中 = 是表格行
    块列表 = 排版.提取块('\n'.join(表格分隔))
    当前小节 = ''
    i = 0
    while i < len(块列表):
        块 = 块列表[i]
        后 = 块列表[i + 1] if i + 1 < len(块列表) else ''
        行 = [x.rstrip() for x in 块.splitlines()]
        if 块.startswith('# '):
            排版.段落(文档, 块[2:], 字号=16, 加粗=True, 居中=True, 行高=22, 段后=6, 跟随=True)
        elif 块.startswith('## '):
            排版.段落(文档, 块[3:], 字号=13, 加粗=True, 行高=18, 段前=5, 段后=3, 跟随=True)
        elif 块.startswith('### '):
            当前小节 = 块[4:]
            排版.段落(文档, 当前小节, 字号=11.5, 加粗=True, 行高=17, 段前=5, 段后=3, 跟随=True)
        elif 块.startswith('**朵朵原答案：**') or 块.startswith('**蓝笔订正：**'):
            排版.段落(文档, 块, 字号=10.5, 行高=14, 段前=2, 段后=1, 跟随=True)
        elif 块 == '**订正：**':
            排版.段落(文档, '**订正：** ' + '_' * 70, 字号=10.5, 行高=18, 段后=2, 跟随=True)
            i += 1
        elif 块 == '**心得：**':
            排版.段落(文档, '**心得：** ' + '_' * 70, 字号=10.5, 行高=18, 段后=6, 跟随=False)
            i += 2
        elif 块.startswith('|'):
            数据 = [[x.strip() for x in y.strip('|').split('|')] for y in 行 if not re.fullmatch(r'[| :\-]+', y)]
            if 数据:
                排版.添加表格(文档, 数据, 宽度=[13, 161] if 数据[0] == ['选项', '句子'] else None)
                排版.段落(文档, '', 行高=4, 段后=0)
        elif len(行) > 1 and all(re.match(r'^[A-D]\. ', x) for x in 行):
            排版.添加选项(文档, 行)
            if 后.startswith('**朵朵原答案：**'):
                文档.paragraphs[-1].paragraph_format.keep_with_next = True
        elif re.match(r'^\*\*\d+\.\*\*', 块):
            # 原稿常把题干和 A-D 选项放在同一个段落中；拆开后才能按短选项横排。
            if len(行) > 1 and all(re.match(r'^[A-D]\. ', x) for x in 行[1:]):
                排版.段落(文档, 行[0], 字号=11, 行高=15, 段前=2, 段后=1, 跟随=True)
                排版.添加选项(文档, 行[1:])
                if 后.startswith(('**朵朵原答案：**', '**蓝笔订正：**')):
                    文档.paragraphs[-1].paragraph_format.keep_with_next = True
                i += 1
                continue
            跟随 = 后.startswith(('**朵朵原答案：**', '**蓝笔订正：**')) or bool(re.match(r'^[A-D]\. ', 后))
            排版.段落(文档, 块, 字号=11, 行高=15, 段前=2, 段后=2, 跟随=跟随)
        elif 块.startswith('＿'):
            排版.段落(文档, re.sub(r'^＿+', '_' * 76, 块), 字号=10.5, 行高=18, 段后=2, 跟随=后.startswith('**'))
        elif len(行) > 1:
            排版.段落(文档, '\n'.join(行), 字号=10.5, 行高=15, 段后=3, 不拆=True)
        else:
            排版.段落(文档, 块, 字号=10.5, 行高=15, 段后=3, 不拆=True)
        i += 1


def 合并(路径列表):
    文档 = 排版.创建文档('全部')
    文档.core_properties.title = '朵朵英语全部错题本'
    文档.core_properties.subject = 'L001、L002、L003、L007、L008、L009、L010 蓝笔订正错题汇总'
    for 序号, 路径 in enumerate(路径列表):
        卷号 = re.search(r'L\d{3}', 路径.name).group()
        if 序号:
            节 = 文档.add_section(WD_SECTION_START.NEW_PAGE)
            节.header.is_linked_to_previous = False
            节.first_page_header.is_linked_to_previous = False
            节.different_first_page_header_footer = False
        else:
            节 = 文档.sections[0]
            节.different_first_page_header_footer = False
        页眉 = 节.header.paragraphs[0]
        页眉.clear()
        排版.加入文字(页眉, f'朵朵英语全部错题本　{卷号}', 9)
        排版正文(文档, 路径.read_text(encoding='utf-8'), 卷号)
    输出 = 输出目录 / '2026-09-17-朵朵英语全部错题本.docx'
    文档.save(输出)
    return 输出


if __name__ == '__main__':
    路径列表 = []
    统计 = {}
    for 卷号 in ('L001', 'L002', 'L003', 'L007', 'L008', 'L009', 'L010'):
        p, n, c = 生成卷Markdown(卷号)
        路径列表.append(p)
        统计[卷号] = {'题数': n, '状态': c, '文件': p.name}
        print(卷号, n, c)
    合并(路径列表)
    (工作目录 / '全部英语蓝笔错题统计.json').write_text(json.dumps(统计, ensure_ascii=False, indent=2), encoding='utf-8')
