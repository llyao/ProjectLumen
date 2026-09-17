from pathlib import Path
import json
import re
import importlib.util
import sys
from collections import Counter
from docx.enum.section import WD_SECTION_START
from docx.shared import Pt, Mm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding='utf-8')
工作目录 = Path(__file__).resolve().parent
根目录 = 工作目录.parent
英语目录 = 根目录 / '朵朵/6年级/记录/英语'
输出目录 = 英语目录 / '错题'
规格 = importlib.util.spec_from_file_location('排版工具', 工作目录/'生成英语空白试卷Word.py')
排版 = importlib.util.module_from_spec(规格)
规格.loader.exec_module(排版)


def 读取核对(卷号):
    数据 = json.loads((工作目录/f'{卷号}错题核对.json').read_text(encoding='utf-8-sig'))
    if 卷号 == 'L008':
        项目 = 数据['基础'][:]
        for 阅读 in 数据['阅读']:
            for 题 in 阅读['错题']:
                项目.append(dict(题, 大题='阅读'+阅读['组']))
    else:
        项目 = 数据
    结果 = {}
    for 题 in 项目:
        大题 = 题['大题']
        阅读字母 = re.search(r'(?:阅读训练\s*|能力提升\s*)([A-D])', 大题)
        if 阅读字母:
            大题 = '阅读'+阅读字母.group(1)
        elif '、' in 大题:
            大题 = 大题.split('、')[0]
        状态 = 题['状态']
        if 卷号 == 'L009' and 大题 == '六' and int(题['题号']) in (10, 11):
            状态 = '未完成'
        if 卷号 == 'L010' and 大题 == '阅读B' and int(题['题号']) in (2, 5):
            # 照片中的最终答案已正确，不作为独立错题。
            continue
        if 状态 not in ('做错', '错题', '未作答', '未完成'):
            continue
        结果[(大题, int(题['题号']))] = dict(题, 状态=状态)
    return 结果


def 分解原稿(卷号):
    文本 = (英语目录/'原稿'/f'2026-09-16-{卷号}英语试题.md').read_text(encoding='utf-8')
    大节 = ''
    小节 = None
    各节 = []
    for 行 in 文本.splitlines():
        if 行.startswith('## '):
            大节 = 行[3:]
            小节 = None
        elif 行.startswith('### '):
            标题 = 行[4:]
            if 大节 == '基础练习':
                键 = 标题.split('、')[0]
            else:
                键 = '阅读'+re.search(r'[A-D]', 标题).group()
            小节 = {'大节': 大节, '标题': 标题, '键': 键, '正文': []}
            各节.append(小节)
        elif 小节 is not None and 行 != '---':
            小节['正文'].append(行)
    for 节 in 各节:
        正文 = '\n'.join(节['正文']).strip()
        if 卷号 == 'L009' and 节['键'] == '一':
            题目 = {}
            for 行 in 正文.splitlines()[2:]:
                for 格 in 行.strip('|').split('|'):
                    号 = int(re.search(r'\*\*(\d+)\.', 格).group(1))
                    题目[号] = 格.strip()
            节['前文'] = ''
            节['题目'] = 题目
        else:
            划分 = list(re.finditer(r'\*\*(\d+)\.\*\*', 正文))
            节['前文'] = 正文[:划分[0].start()].strip() if 划分 else 正文
            节['题目'] = {int(配对.group(1)): 正文[配对.start():划分[号+1].start() if 号+1<len(划分) else len(正文)].strip()
                             for 号,配对 in enumerate(划分)}
    return 各节


def 反馈(题, 选择文本=''):
    原答案 = 题['原答案']
    if len(原答案) == 1 and 原答案 in 'ABCD' and 选择文本:
        匹配 = re.search(r'^'+原答案+r'\. (.+)', 选择文本, re.M)
        if 匹配:
            原答案 += '. '+匹配.group(1).strip()
    if 题['状态'] == '未完成' and not 原答案.startswith('未完成'):
        原答案 = '未完成；'+原答案
    return '\n\n'.join([
        '**朵朵原答案：** '+原答案,
        '**订正：**', '＿＿'*24,
        '**心得：**', '＿＿'*24, '＿＿'*24,
    ])


def 生成Markdown(卷号):
    错题 = 读取核对(卷号)
    各节 = 分解原稿(卷号)
    组成 = [f'# 六年级第一学期英语练习 {卷号}']
    当前大节 = ''
    统计 = []
    for 节 in 各节:
        本节错题 = {号:题 for (键,号),题 in 错题.items() if 键 == 节['键']}
        if not 本节错题:
            continue
        if 当前大节 != 节['大节']:
            组成.append('## '+节['大节'])
            当前大节 = 节['大节']
        组成.append('### '+节['标题'])
        阅读 = 节['键'].startswith('阅读')
        连续对话 = 卷号 == 'L010' and 节['键'] == '三'
        if 节['前文']:
            组成.append(节['前文'])
        if not 节['题目']:
            # 首字母阅读的题目都在文章中；下方沿用原空号留出订正位置。
            if 连续对话:
                组成 += ['**5.** ＿＿＿＿＿＿＿＿', 反馈(本节错题[5])]
            else:
                for 号 in sorted(本节错题):
                    组成.append(f'**{号}.** ＿＿＿＿＿＿＿＿')
                    组成.append(反馈(本节错题[号]))
            统计.append({'大题':节['键'], '收录题号':list(range(1,7)) if 连续对话 else list(range(1,6)), '订正题号':list(本节错题)})
            continue
        收录号 = list(节['题目']) if 阅读 else list(本节错题)
        for 号 in 收录号:
            原题 = 节['题目'][号]
            组成.append(原题)
            if 号 in 本节错题:
                组成.append(反馈(本节错题[号], 原题))
        统计.append({'大题':节['键'], '收录题号':收录号, '订正题号':list(本节错题)})
    路径 = 输出目录/f'2026-09-16-{卷号}英语错题集.md'
    路径.write_text('\n\n'.join(组成)+'\n', encoding='utf-8')
    print(卷号, '需订正题数', len(错题), '分类', dict(Counter(题['状态'] for 题 in 错题.values())))
    return 路径, 统计


def 排版正文(文档, 文本, 卷号):
    块列表 = 排版.提取块(文本)
    当前大节 = ''
    当前小节 = ''
    索引 = 0
    while 索引 < len(块列表):
        块 = 块列表[索引]
        后块 = 块列表[索引+1] if 索引+1<len(块列表) else ''
        各行 = [行.rstrip() for 行 in 块.splitlines()]
        if 块.startswith('# '):
            排版.段落(文档, 块[2:], 字号=17, 加粗=True, 居中=True, 行高=25, 段后=10, 跟随=True)
        elif 块.startswith('## '):
            当前大节 = 块[3:]
            行 = 排版.段落(文档, 当前大节, 字号=14, 加粗=True, 行高=20, 段前=8, 段后=5, 跟随=True)
        elif 块.startswith('### '):
            当前小节 = 块[4:]
            行 = 排版.段落(文档, 当前小节, 字号=12.5, 加粗=True, 行高=19, 段前=7, 段后=5, 跟随=True)
        elif 块.startswith('**朵朵原答案：**'):
            排版.段落(文档, 块, 字号=11, 行高=16, 段前=4, 段后=4, 跟随=True)
        elif 块 == '**订正：**':
            行 = 排版.段落(文档, '**订正：** '+ '_'*70, 字号=11, 行高=24, 段后=5, 跟随=True)
            索引 += 1
        elif 块 == '**心得：**':
            排版.段落(文档, '**心得：** '+ '_'*70, 字号=11, 行高=24, 段后=0, 跟随=True)
            排版.段落(文档, '_'*78, 字号=11, 行高=24, 段后=12, 跟随=False)
            索引 += 2
        elif 块.startswith('|'):
            数据 = [[格.strip() for 格 in 行.strip('|').split('|')] for 行 in 各行 if not re.fullmatch(r'[| :\-]+', 行)]
            排版.添加表格(文档, 数据, 宽度=[13,161] if 数据[0]==['选项','句子'] else None)
            排版.段落(文档, '', 行高=4, 段后=0)
        elif len(各行) > 1 and all(re.match(r'^[A-D]\. ', 行) for 行 in 各行):
            排版.添加选项(文档, 各行)
            if 后块.startswith('**朵朵原答案：**'):
                文档.paragraphs[-1].paragraph_format.keep_with_next = True
        elif 卷号 == 'L008' and 当前小节.startswith('一、') and len(各行)==8:
            排版.添加表格(文档, [各行[:4],各行[4:]], 表头=False, 边框=False)
            for 行 in 文档.tables[-1].rows[-1].cells:
                行.paragraphs[0].paragraph_format.keep_with_next=True
        elif re.match(r'^\*\*\d+\.\*\*', 块):
            if re.fullmatch(r'\*\*\d+\.\*\*\s*（\s*）',块) and re.match(r'^[A-D]\. ',后块):
                排版.添加选项(文档, [行.rstrip() for 行 in 后块.splitlines()], 题号=块)
                文档.paragraphs[-1].paragraph_format.keep_with_next = (索引+2<len(块列表) and 块列表[索引+2].startswith('**朵朵原答案：**'))
                索引 += 2
                continue
            跟随 = 后块.startswith('**朵朵原答案：**') or bool(re.match(r'^[A-D]\. |^＿',后块)) or (卷号=='L008' and 当前小节.startswith('一、'))
            排版.段落(文档, 块, 行高=17, 段前=3, 段后=4, 跟随=跟随)
        elif 块.startswith('＿'):
            if re.fullmatch(r'＿+[!！?？]?',块) and 后块.startswith('**朵朵原答案：**'):
                # 单独的空白作答横线由下面的订正书写区承接，保留原有句末标点。
                if 块[-1] in '!！?？':
                    文档.paragraphs[-1].add_run(块[-1])
                索引 += 1
                continue
            if re.fullmatch(r'＿+[!！?？]?',块):
                块 = re.sub(r'^＿+', '_'*76, 块)
            排版.段落(文档, 块, 行高=22, 段后=4, 跟随=后块.startswith('**朵朵原答案：**'))
        elif len(各行)>1:
            排版.段落(文档, '\n'.join(各行), 行高=17, 段后=5, 不拆=True)
        else:
            是说明=块.startswith('**') and 块.endswith('**')
            排版.段落(文档, 块, 行高=17, 段后=6, 跟随=是说明, 不拆=True)
        索引 += 1


def 合并Word(路径列表):
    文档 = 排版.创建文档('L008')
    文档.core_properties.title = '朵朵英语错题集'
    文档.core_properties.subject = 'L008、L009、L010 按试卷连续整理'
    for 序号, 路径 in enumerate(路径列表):
        卷号 = re.search(r'L\d{3}',路径.name).group()
        if 序号:
            节 = 文档.add_section(WD_SECTION_START.NEW_PAGE)
            节.header.is_linked_to_previous = False
            节.first_page_header.is_linked_to_previous = False
            节.different_first_page_header_footer = False
        else:
            节 = 文档.sections[0]
            节.different_first_page_header_footer = False
        页眉 = 节.header.paragraphs[0]
        页眉.clear()
        排版.加入文字(页眉, f'朵朵英语错题集　{卷号}', 9)
        排版正文(文档, 路径.read_text(encoding='utf-8'), 卷号)
    输出 = 输出目录/'2026-09-16-朵朵英语错题集-L008至L010.docx'
    文档.save(输出)
    print(输出)


if __name__ == '__main__':
    所有路径 = []
    核对统计 = {}
    for 卷号 in ['L008','L009','L010']:
        路径, 统计 = 生成Markdown(卷号)
        所有路径.append(路径)
        核对统计[卷号] = 统计
    (工作目录/'错题集收录清单.json').write_text(json.dumps(核对统计,ensure_ascii=False,indent=2),encoding='utf-8')
    合并Word(所有路径)
