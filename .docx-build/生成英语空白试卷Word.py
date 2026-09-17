from pathlib import Path
import re
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


根目录 = Path(__file__).resolve().parent.parent
原稿目录 = 根目录 / '朵朵/6年级/记录/英语/原稿'


def 设置字体(运行, 字号=12, 加粗=False, 下划线=False):
    运行.font.name = 'Times New Roman'
    运行.font.size = Pt(字号)
    运行.font.bold = 加粗
    运行.font.underline = 下划线
    运行.font.color.rgb = RGBColor(0, 0, 0)
    字体 = 运行._element.get_or_add_rPr().get_or_add_rFonts()
    字体.set(qn('w:eastAsia'), '宋体')
    字体.set(qn('w:ascii'), 'Times New Roman')
    字体.set(qn('w:hAnsi'), 'Times New Roman')


def 加入文字(段落, 文本, 字号=12, 加粗=False):
    # 原稿的组合下划线转为 Word 原生下划线；答题横线保持空白。
    for 序号, 片段 in enumerate(re.split(r'\*\*(.*?)\*\*', 文本)):
        # 使用连续英文横线，防止首字母填空在横线中间被自动换行。
        片段 = 片段.replace('＿', '__')
        当前加粗 = 加粗 or 序号 % 2 == 1
        索引 = 0
        缓冲 = ''
        当前下划线 = False
        while 索引 < len(片段):
            字符 = 片段[索引]
            if 字符 == '\u0332':
                索引 += 1
                continue
            下划线 = 索引 + 1 < len(片段) and 片段[索引 + 1] == '\u0332'
            if 缓冲 and 下划线 != 当前下划线:
                设置字体(段落.add_run(缓冲), 字号, 当前加粗, 当前下划线)
                缓冲 = ''
            当前下划线 = 下划线
            缓冲 += 字符
            索引 += 2 if 下划线 else 1
        if 缓冲:
            设置字体(段落.add_run(缓冲), 字号, 当前加粗, 当前下划线)


def 设置段落(段落, 段前=0, 段后=4, 行高=16, 跟随=False, 不拆=True):
    格式 = 段落.paragraph_format
    格式.space_before = Pt(段前)
    格式.space_after = Pt(段后)
    格式.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    格式.line_spacing = Pt(行高)
    格式.keep_with_next = 跟随
    格式.keep_together = 不拆
    格式.widow_control = True
    return 段落


def 段落(文档, 文本='', **选项):
    字号 = 选项.pop('字号', 12)
    加粗 = 选项.pop('加粗', False)
    居中 = 选项.pop('居中', False)
    结果 = 设置段落(文档.add_paragraph(), **选项)
    if 居中:
        结果.alignment = WD_ALIGN_PARAGRAPH.CENTER
    加入文字(结果, 文本, 字号, 加粗)
    return 结果


def 加入页码(页脚):
    行 = 页脚.paragraphs[0]
    行.alignment = WD_ALIGN_PARAGRAPH.CENTER
    设置段落(行, 段后=0, 行高=12)
    for 项目 in ['第 ', 'PAGE', ' 页 / 共 ', 'NUMPAGES', ' 页']:
        运行 = 行.add_run()
        设置字体(运行, 9)
        if 项目 in ('PAGE', 'NUMPAGES'):
            域 = OxmlElement('w:fldSimple')
            域.set(qn('w:instr'), 项目)
            运行._r.addnext(域)
        else:
            运行.text = 项目


def 创建文档(卷号):
    文档 = Document()
    节 = 文档.sections[0]
    节.page_width, 节.page_height = Mm(210), Mm(297)
    节.top_margin, 节.bottom_margin = Mm(17), Mm(17)
    节.left_margin, 节.right_margin = Mm(18), Mm(18)
    节.header_distance, 节.footer_distance = Mm(8), Mm(8)
    节.different_first_page_header_footer = True
    页眉 = 节.header.paragraphs[0]
    页眉.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    设置段落(页眉, 行高=12, 段后=0)
    加入文字(页眉, f'六年级英语练习 {卷号}', 9)
    加入页码(节.footer)
    加入页码(节.first_page_footer)
    样式 = 文档.styles['Normal']
    样式.font.name = 'Times New Roman'
    样式.font.size = Pt(12)
    样式.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')
    样式.paragraph_format.line_spacing = Pt(16)
    样式.paragraph_format.space_after = Pt(4)
    文档.core_properties.title = f'六年级第一学期英语练习 {卷号}'
    文档.core_properties.subject = '英语试题'
    文档.core_properties.author = ''
    文档.core_properties.last_modified_by = ''
    文档.core_properties.comments = ''
    文档.core_properties.revision = 1
    兼容 = 文档.settings.element.find(qn('w:compat'))
    if 兼容 is not None:
        for 子项 in list(兼容):
            if 子项.tag == qn('w:compatSetting') and 子项.get(qn('w:name')) == 'compatibilityMode':
                子项.set(qn('w:val'), '15')
    return 文档


def 添加表格(文档, 各行, 表头=True, 边框=True, 字号=11.5, 宽度=None):
    表 = 文档.add_table(rows=0, cols=len(各行[0]))
    表.alignment = WD_TABLE_ALIGNMENT.CENTER
    表.autofit = False
    总宽 = 174
    if 宽度 is None:
        宽度 = [总宽 / len(各行[0])] * len(各行[0])
    for 列, 毫米 in zip(表.columns, 宽度):
        列.width = Mm(毫米)
    边界 = OxmlElement('w:tblBorders')
    for 名称 in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        属性 = OxmlElement('w:' + 名称)
        属性.set(qn('w:val'), 'single' if 边框 else 'nil')
        属性.set(qn('w:sz'), '4')
        属性.set(qn('w:color'), '000000')
        边界.append(属性)
    表._tbl.tblPr.append(边界)
    for 行号, 各格 in enumerate(各行):
        行 = 表.add_row()
        行属性 = 行._tr.get_or_add_trPr()
        行属性.append(OxmlElement('w:cantSplit'))
        if 行号 == 0 and 表头:
            行属性.append(OxmlElement('w:tblHeader'))
        for 格号, (格, 内容) in enumerate(zip(行.cells, 各格)):
            格.width = Mm(宽度[格号])
            格.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            格属性 = 格._tc.get_or_add_tcPr()
            内边距 = OxmlElement('w:tcMar')
            for 方位, 数值 in [('top', 40), ('bottom', 40), ('left', 75), ('right', 75)]:
                项 = OxmlElement('w:' + 方位)
                项.set(qn('w:w'), str(数值))
                项.set(qn('w:type'), 'dxa')
                内边距.append(项)
            格属性.append(内边距)
            行段 = 格.paragraphs[0]
            设置段落(行段, 行高=16, 段后=0, 跟随=行号 < len(各行)-1)
            加入文字(行段, 内容, 字号, 表头 and 行号 == 0)
    return 表


def 添加选项(文档, 各项, 题号=''):
    最长 = max(len(项.replace('\u0332', '')) for 项 in 各项)
    if len(各项) == 2:
        列数 = 2 if 最长 <= 50 else 1
    elif len(各项) == 3:
        列数 = 3 if 最长 <= 26 else 1
    else:
        列数 = 4 if 最长 <= 18 else (2 if 最长 <= 38 else 1)
    for 起点 in range(0, len(各项), 列数):
        末行 = 起点 + 列数 >= len(各项)
        行 = 设置段落(文档.add_paragraph(), 段后=5 if 末行 else 1, 跟随=not 末行)
        起始位置 = 26 if 题号 else 0
        for 序号 in range(0 if 题号 else 1, 列数):
            行.paragraph_format.tab_stops.add_tab_stop(Mm(起始位置 + (174-起始位置) / 列数 * 序号))
        加入文字(行, (题号+'\t' if 题号 else '') + '\t'.join(各项[起点:起点+列数]), 11.5)


def 提取块(文本):
    return [块.strip() for 块 in re.split(r'\n\s*\n', 文本.replace('\r\n', '\n')) if 块.strip()]


def 转换(输入):
    卷号 = re.search(r'L\d{3}', 输入.name).group()
    文档 = 创建文档(卷号)
    各块 = 提取块(输入.read_text(encoding='utf-8'))
    当前小节 = ''
    索引 = 0
    while 索引 < len(各块):
        块 = 各块[索引]
        后块 = 各块[索引+1] if 索引+1 < len(各块) else ''
        各行 = [行.rstrip() for 行 in 块.splitlines()]
        if 块 == '---':
            pass
        elif 块.startswith('# '):
            段落(文档, 块[2:], 字号=17, 加粗=True, 居中=True, 行高=24, 段后=5, 跟随=True)
        elif 块.startswith('## '):
            段落(文档, 块[3:], 字号=13, 加粗=True, 行高=19, 段前=6, 段后=3, 跟随=True)
        elif 块.startswith('### '):
            当前小节 = 块[4:]
            段落(文档, 当前小节, 字号=12, 加粗=True, 行高=17, 段前=6, 段后=4, 跟随=True)
        elif 块.startswith('**Exercise for Grade'):
            段落(文档, 块, 字号=11.5, 居中=True, 行高=17, 段后=6, 跟随=True)
        elif 块.startswith('班级：'):
            for 行号, 行 in enumerate(各行):
                段落(文档, 行, 字号=10.5, 行高=17, 段后=1, 跟随=True)
        elif 块.startswith('|'):
            数据 = [[格.strip() for 格 in 行.strip('|').split('|')] for 行 in 各行 if not re.fullmatch(r'[| :\-]+', 行)]
            宽度 = [13, 161] if 数据[0] == ['选项', '句子'] else None
            添加表格(文档, 数据, 宽度=宽度)
            段落(文档, '', 行高=3, 段后=0)
        elif len(各行) > 1 and all(re.match(r'^[A-D]\. ', 行) for 行 in 各行):
            添加选项(文档, 各行)
        elif 卷号 == 'L008' and 当前小节.startswith('一、') and len(各行) == 8:
            添加表格(文档, [各行[:4], 各行[4:]], 表头=False, 边框=False)
        elif re.match(r'^\*\*\d+\.\*\*', 块):
            是单题 = len(re.findall(r'\*\*\d+\.\*\*', 块)) == 1
            后是作答 = bool(re.match(r'^＿+', 后块))
            后是选项 = bool(re.match(r'^[A-D]\. ', 后块))
            后是提示 = 后块.startswith("Let's go") or 后块.startswith('He has 8 lessons')
            后是辨音 = 卷号 == 'L008' and 当前小节.startswith('一、')
            if re.fullmatch(r'\*\*\d+\.\*\*\s*（\s*）', 块) and 后是选项:
                添加选项(文档, [行.rstrip() for 行 in 后块.splitlines()], 题号=块)
                索引 += 1
            elif 当前小节.startswith('三、翻译') and 后是作答 and len(块) < 26:
                段落(文档, 块 + '　' + '＿'*18, 段后=6, 行高=21)
                索引 += 1
            else:
                行 = 段落(文档, 块.replace('  \n', '\n'), 段后=2 if 后是作答 or 后是选项 or 后是提示 else 5,
                         跟随=是单题 and (后是作答 or 后是选项 or 后是提示 or 后是辨音))
        elif re.match(r'^＿+', 块):
            答题文字 = re.sub(r'^＿+', '＿'*38, 块) if re.fullmatch(r'＿+[!！?？]?', 块) else 块
            段落(文档, 答题文字, 行高=24, 段后=6)
        elif 块.startswith("Let's go") or 块.startswith('He has 8 lessons'):
            段落(文档, 块, 行高=22, 段后=6)
        elif len(各行) > 1:
            行程 = 块.startswith(('Tuesday:', 'Wednesday:', 'Thursday:'))
            for 行号, 行 in enumerate(各行):
                段落(文档, 行, 段后=3, 不拆=True, 跟随=行程 and 行号 < len(各行)-1)
        else:
            是题目说明 = 块.startswith('**') and 块.endswith('**')
            段落(文档, 块, 段后=5, 跟随=是题目说明, 不拆=len(块) < 250)
        索引 += 1
    # 在题组边界主动分页，避免尾页只留下少量题目。
    分页标题 = {
        'L008': ['四、选出正确的答案。', '六、用所给单词的适当形式填空', '阅读训练 B'],
        'L009': ['四、选出正确的答案', '六、改写句子'],
        'L010': ['三、选出正确的句子补全对话。', '五、用括号中所给单词的适当形式完成下列句子。',
                 'B. Choose the best words or expressions to complete the passage.', '★ 能力提升'],
    }
    for 行 in 文档.paragraphs:
        格式 = 行.paragraph_format
        if 卷号 != 'L010':
            if 格式.line_spacing is not None and isinstance(格式.line_spacing, int):
                高度 = 格式.line_spacing.pt
                是答题行 = 高度 == 24 and 行.text.startswith(('＿', '_'))
                格式.line_spacing = Pt((20.5 if 卷号 == 'L008' else 22) if 是答题行 else 高度 * 0.94)
            if 格式.space_after is not None:
                格式.space_after = Pt(max(0, 格式.space_after.pt-1))
        if 行.text in 分页标题[卷号]:
            格式.page_break_before = True
    输出 = 输入.with_suffix('.docx')
    文档.save(输出)
    print(输出)


if __name__ == '__main__':
    for 路径 in sorted(原稿目录.glob('2026-09-16-L0*英语试题.md')):
        转换(路径)
